# -*- coding: utf-8 -*-
"""
把 docs/research_report.md 转换为 Word 文档（供作品集使用）。

针对本报告使用的 Markdown 子集实现解析：
  # / ## / ### 标题、表格、图片 ![alt](path)、**加粗**、有序/无序列表、--- 分隔线、普通段落

用法：
    python src/md_report_to_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "docs" / "research_report.md"
OUT_PATH = ROOT / "docs" / "AI浪潮下的技能需求变迁_研究报告.docx"


def 设置中文字体(doc: Document) -> None:
    """正文宋体小四、标题微软雅黑，保证中文显示。"""
    正文 = doc.styles["Normal"]
    正文.font.name = "Times New Roman"
    正文.font.size = Pt(12)
    正文.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for 级别, 字号 in [("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)]:
        样式 = doc.styles[级别]
        样式.font.name = "Arial"
        样式.font.size = Pt(字号)
        样式.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
        样式.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def 写入加粗混合文本(段落, 文本: str) -> None:
    """按 **加粗** 标记拆分写入 run。"""
    for 片段 in re.split(r"(\*\*[^*]+\*\*)", 文本):
        if not 片段:
            continue
        if 片段.startswith("**") and 片段.endswith("**"):
            段落.add_run(片段[2:-2]).bold = True
        else:
            段落.add_run(片段)


def 转换表格(doc: Document, 行列表: list[str]) -> None:
    """把 Markdown 表格行（含分隔行）转成 Word 表格。"""
    行数据 = []
    for 行 in 行列表:
        if re.match(r"^\s*\|[\s\-|]+\|\s*$", 行):  # 分隔行 |---|---|
            continue
        单元 = [c.strip() for c in 行.strip().strip("|").split("|")]
        行数据.append(单元)
    if not 行数据:
        return
    表格 = doc.add_table(rows=len(行数据), cols=len(行数据[0]))
    表格.style = "Light Grid Accent 1"
    for i, 单元行 in enumerate(行数据):
        for j, 值 in enumerate(单元行):
            格 = 表格.cell(i, j)
            格.text = ""
            写入加粗混合文本(格.paragraphs[0], 值)
            for r in 格.paragraphs[0].runs:
                r.font.size = Pt(10)
                if i == 0:
                    r.bold = True


def main() -> None:
    文本 = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    设置中文字体(doc)

    行缓存: list[str] = []  # 表格行缓存

    def 冲刷表格() -> None:
        nonlocal 行缓存
        if 行缓存:
            转换表格(doc, 行缓存)
            行缓存 = []

    for 行 in 文本.splitlines():
        if 行.strip().startswith("|"):
            行缓存.append(行)
            continue
        冲刷表格()

        if not 行.strip():
            continue
        if 行.strip() == "---":
            continue  # 分隔线在 Word 中以留白代替
        m = re.match(r"^(#{1,3})\s+(.*)$", 行)
        if m:
            级别 = len(m.group(1))
            doc.add_heading(m.group(2).replace("**", ""), level=级别)
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", 行.strip())
        if m:
            图路径 = (MD_PATH.parent / m.group(2)).resolve()
            if 图路径.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(图路径), width=Cm(15.5))
                题注 = doc.add_paragraph(m.group(1))
                题注.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in 题注.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                doc.add_paragraph(f"[图片缺失：{m.group(2)}]")
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", 行.strip())
        if m:
            写入加粗混合文本(doc.add_paragraph(style="List Number"), m.group(2))
            continue
        if 行.strip().startswith("- "):
            写入加粗混合文本(doc.add_paragraph(style="List Bullet"), 行.strip()[2:])
            continue
        if 行.lstrip().startswith("> "):
            p = doc.add_paragraph()
            写入加粗混合文本(p, 行.lstrip()[2:])
            p.paragraph_format.left_indent = Cm(0.8)
            continue
        # 普通段落（跳过 Markdown 代码围栏标记，保留代码文本）
        if 行.strip().startswith("```"):
            continue
        p = doc.add_paragraph()
        写入加粗混合文本(p, 行)

    冲刷表格()

    # 页脚页码说明（简化：加一行生成信息）
    节 = doc.sections[0]
    节.footer.paragraphs[0].text = "skill-shift-analysis 项目研究报告 · 由可复现代码生成"
    节.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(OUT_PATH)
    print(f"已生成: {OUT_PATH} ({OUT_PATH.stat().st_size} 字节)")


if __name__ == "__main__":
    main()
